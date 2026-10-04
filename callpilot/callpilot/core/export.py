"""Save a call to the computer: transcript (TXT, DOCX, SRT), summary, pins, recording."""

from __future__ import annotations

import csv
import datetime as dt
import json
import re
import shutil
from pathlib import Path

from callpilot.core.sessions import export_srt, export_text, field_label, redact_record


def default_export_dir() -> Path:
    home = Path.home()
    docs = home / "Documents"
    base = (docs if docs.exists() else home) / "CallPilot" / "Calls"
    base.mkdir(parents=True, exist_ok=True)
    return base


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip() or "call"


def call_folder_name(rec: dict) -> str:
    stamp = dt.datetime.fromtimestamp(rec.get("started_at", 0)).strftime("%Y-%m-%d %H-%M-%S")
    who = (rec.get("fields") or {}).get("caller_name") or (rec.get("fields") or {}).get("client_name") or \
        (rec.get("fields") or {}).get("handler_name") or ""
    hub = rec.get("hub_name", "call").split("–")[0].strip()
    return _safe(f"{stamp} {hub}" + (f" – {who}" if who else ""))


def transcript_docx(rec: dict, path: Path) -> Path:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.add_heading(f"Call transcript – {rec.get('hub_name', '')}", level=1)
    started = dt.datetime.fromtimestamp(rec.get("started_at", 0))
    meta = [f"Date: {started:%A %d %B %Y, %H:%M}", f"Call id: {rec.get('call_id', '')}",
            f"Call type: {rec.get('call_type', '') or '–'}"]
    if rec.get("notice_given_at"):
        meta.append(f"Recording notice given at {dt.datetime.fromtimestamp(rec['notice_given_at']):%H:%M:%S}")
    d.add_paragraph("\n".join(meta))
    summ = rec.get("summary") or {}
    if summ.get("summary"):
        d.add_heading("Summary", level=2)
        d.add_paragraph(summ["summary"])
    if summ.get("file_note"):
        d.add_heading("File note", level=2)
        d.add_paragraph(summ["file_note"])
    d.add_heading("Transcript", level=2)
    for s in rec.get("segments", []):
        t = dt.datetime.fromtimestamp(s["start"]).strftime("%H:%M:%S")
        who = "Us" if s["speaker"] == "agent" else "Caller"
        p = d.add_paragraph()
        r = p.add_run(f"[{t}] {who}: ")
        r.bold = True
        p.add_run(s["text"])
        if s.get("translation"):
            tr = d.add_paragraph("    ↳ " + s["translation"])
            tr.runs[0].font.size = Pt(9)
    pins = rec.get("pins") or []
    if pins:
        d.add_heading("Pins (dates, commitments, figures)", level=2)
        for p in pins:
            d.add_paragraph(f"{p.get('kind', '')}: {p.get('value', '')} – “{p.get('quote', '')}”"
                            + (f" (due {p['due_at']})" if p.get("due_at") else ""), style="List Bullet")
    fields = {k: v for k, v in (rec.get("fields") or {}).items() if v}
    if fields:
        d.add_heading("Details captured", level=2)
        labels = rec.get("field_labels") or {}
        for k, v in fields.items():
            d.add_paragraph(f"{field_label(k, labels)}: {v}", style="List Bullet")
    d.save(str(path))
    return path


def export_call(rec: dict, dest: Path | None = None, vault=None, include_recording: bool = True) -> Path:
    """Write everything about one call into a folder and return it."""
    base = dest or default_export_dir()
    folder = base / call_folder_name(rec)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "transcript.txt").write_text(export_text(rec), encoding="utf-8")
    srt = export_srt(rec)
    if srt:
        (folder / "transcript.srt").write_text(srt, encoding="utf-8")
    try:
        transcript_docx(rec, folder / "transcript.docx")
    except Exception:  # noqa: BLE001 - docx is optional
        pass
    (folder / "call.json").write_text(json.dumps(rec, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    pins = rec.get("pins") or []
    if pins:
        with (folder / "pins.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["kind", "value", "quote", "speaker", "due", "time"])
            for p in pins:
                w.writerow([p.get("kind"), p.get("value"), p.get("quote"), p.get("speaker"), p.get("due_at"),
                            dt.datetime.fromtimestamp(p.get("at", 0)).strftime("%H:%M:%S")])
    email = ((rec.get("wrapup") or {}).get("email")) or ((rec.get("summary") or {}).get("email")) or {}
    if email.get("body"):
        (folder / "as-discussed-email.txt").write_text(
            f"To: {email.get('to', '')}\nSubject: {email.get('subject', '')}\n\n{email['body']}", encoding="utf-8")
    if include_recording and rec.get("recording"):
        src = Path(rec["recording"])
        try:
            if src.suffix == ".cpv" and vault is not None and src.exists():
                (folder / "recording.wav").write_bytes(vault.decrypt(src.read_bytes()))
            elif src.exists():
                shutil.copy2(src, folder / "recording.wav")
        except Exception:  # noqa: BLE001
            pass
    return folder


def export_dir_for(settings) -> Path:
    return Path(settings.export.folder) if settings.export.folder else default_export_dir()


def auto_export(rec: dict, settings, vault=None) -> Path | None:
    """The after-every-call save, honouring the privacy settings: nothing is written when
    sessions are not kept, PII is redacted when the store redacts, and the recording only
    goes out when it was kept. Returns the folder, or None when nothing was written."""
    priv = settings.privacy
    if not settings.export.auto_save_calls or not priv.save_sessions:
        return None
    if priv.redact_saved_pii:
        rec = redact_record(rec)
    return export_call(rec, export_dir_for(settings), vault=vault, include_recording=bool(rec.get("recording")))
